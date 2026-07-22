import os
import pandas as pd
from typing import List
from pypdf import PdfReader
from pptx import Presentation
from langchain_core.documents import Document

try:
    import docx2txt
except ImportError:
    docx2txt = None

def extract_documents_from_file(file_path: str) -> List[Document]:
    """
    Extracts text from PDF, PPTX, DOCX, CSV, and Excel files while attaching page, slide, or row metadata.
    """
    documents = []
    file_name = os.path.basename(file_path)
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        reader = PdfReader(file_path)
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                documents.append(
                    Document(
                        page_content=page_text.strip(),
                        metadata={"page": idx + 1, "source": file_name, "type": "pdf"}
                    )
                )

    elif ext == ".pptx":
        prs = Presentation(file_path)
        for idx, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                full_slide_text = "\n".join(slide_texts)
                documents.append(
                    Document(
                        page_content=full_slide_text,
                        metadata={"slide": idx + 1, "source": file_name, "type": "pptx"}
                    )
                )

    elif ext in [".docx", ".doc"]:
        if docx2txt:
            raw_text = docx2txt.process(file_path)
            if raw_text and raw_text.strip():
                # Split by double newlines for paragraph sections
                paragraphs = [p.strip() for p in raw_text.split("\n\n") if p.strip()]
                for idx, para in enumerate(paragraphs):
                    documents.append(
                        Document(
                            page_content=para,
                            metadata={"section": f"Paragraph {idx+1}", "source": file_name, "type": "docx"}
                        )
                    )

    elif ext in [".txt", ".text", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            raw_text = f.read()
        if raw_text and raw_text.strip():
            paragraphs = [p.strip() for p in raw_text.split("\n\n") if p.strip()]
            for idx, para in enumerate(paragraphs):
                documents.append(
                    Document(
                        page_content=para,
                        metadata={"section": f"Section {idx+1}", "source": file_name, "type": "txt"}
                    )
                )

    elif ext in [".csv", ".xlsx"]:
        if ext == ".csv":
            df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)

        headers = ", ".join(df.columns.tolist())
        summary_text = f"Dataset Summary:\nColumns: {headers}\nTotal Rows: {len(df)}\n"
        summary_text += f"Sample Data:\n{df.head(10).to_string()}\n"
        documents.append(
            Document(
                page_content=summary_text,
                metadata={"section": "Summary & Overview", "source": file_name, "type": "dataframe"}
            )
        )

        chunk_size = 25
        for i in range(0, len(df), chunk_size):
            chunk_df = df.iloc[i : i + chunk_size]
            documents.append(
                Document(
                    page_content=f"Columns: {headers}\nRows {i+1} to {min(i+chunk_size, len(df))}:\n{chunk_df.to_string()}",
                    metadata={"rows": f"{i+1}-{min(i+chunk_size, len(df))}", "source": file_name, "type": "dataframe"}
                )
            )

    return documents
